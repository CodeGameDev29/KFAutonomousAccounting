"""core/ledger.py — the ledger workbook and the CSV exports written from it.

Every figure and merchant below is invented. What is pinned is the shape an
accountant (or a spreadsheet) opening the output depends on: month labels and
per-account sheet names, the standard header row, the IndexSheet and NET INCOME
sheets an empty book still has to carry, a receipt link that stays clickable,
appends that add rows instead of overwriting them, and the transactions CSV's
date and column layout.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from core.ledger import (
    LEDGER_HEADERS_STANDARD,
    export_ledger,
    export_ledger_csv,
    export_transactions_csv,
    month_label,
    sheet_name,
    write_ledger_entries,
)
from models.ledger_entry import LedgerEntry
from models.transaction import AccountType, Transaction, TransactionStatus

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_entry(
    account: AccountType = AccountType.CAD,
    month: str = "Jan2026",
    date_posted: date = date(2026, 1, 15),
    amount: Decimal = Decimal("-49.99"),
    currency: str = "CAD",
    description: str = "Example Retailer Office Supplies",
    category: str = "Software and IT",
    document_link: str | None = None,
    note: str | None = None,
    transaction_type: str = "DEBIT",
) -> LedgerEntry:
    return LedgerEntry(
        account=account,
        month=month,
        transaction_type=transaction_type,
        date_posted=date_posted,
        amount=amount,
        currency=currency,
        description=description,
        category=category,
        document_link=document_link,
        note=note,
    )


# ---------------------------------------------------------------------------
# month_label
# ---------------------------------------------------------------------------

class TestMonthLabel:
    def test_month_label(self):
        assert month_label(2026, 1) == "Jan2026"

    def test_month_label_december(self):
        assert month_label(2026, 12) == "Dec2026"


# ---------------------------------------------------------------------------
# sheet_name
# ---------------------------------------------------------------------------

class TestSheetName:
    def test_sheet_name_cad(self):
        assert sheet_name(AccountType.CAD, 2026, 1) == "Jan2026_CAD"

    def test_sheet_name_credit_card(self):
        assert sheet_name(AccountType.CREDIT_CARD, 2026, 3) == "Mar2026_CreditCard"


# ---------------------------------------------------------------------------
# export_ledger — file creation
# ---------------------------------------------------------------------------

class TestExportLedgerCreatesFile:
    def test_export_ledger_creates_file(self, tmp_path: Path):
        out = tmp_path / "ledger.xlsx"
        entries = {
            "Jan2026_CAD": [_make_entry()],
        }
        export_ledger(entries, out, year=2026)
        assert out.exists()


# ---------------------------------------------------------------------------
# export_ledger — monthly sheets
# ---------------------------------------------------------------------------

class TestExportLedgerMonthlySheets:
    def test_export_ledger_creates_monthly_sheets(self, tmp_path: Path):
        out = tmp_path / "ledger.xlsx"
        entries = {
            "Jan2026_CAD": [_make_entry(month="Jan2026", date_posted=date(2026, 1, 10))],
            "Feb2026_CAD": [
                _make_entry(
                    month="Feb2026",
                    date_posted=date(2026, 2, 5),
                    description="Feb purchase",
                )
            ],
        }
        export_ledger(entries, out, year=2026)
        wb = openpyxl.load_workbook(out)
        assert "Jan2026_CAD" in wb.sheetnames
        assert "Feb2026_CAD" in wb.sheetnames
        wb.close()


# ---------------------------------------------------------------------------
# export_ledger — headers
# ---------------------------------------------------------------------------

class TestExportLedgerHeaders:
    def test_export_ledger_standard_headers(self, tmp_path: Path):
        out = tmp_path / "ledger.xlsx"
        entries = {
            "Jan2026_CAD": [_make_entry()],
        }
        export_ledger(entries, out, year=2026)
        wb = openpyxl.load_workbook(out)
        ws = wb["Jan2026_CAD"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, 9)]
        assert headers == LEDGER_HEADERS_STANDARD
        wb.close()


# ---------------------------------------------------------------------------
# export_ledger — entry data
# ---------------------------------------------------------------------------

class TestExportLedgerEntryData:
    def test_export_ledger_entry_data(self, tmp_path: Path):
        entry = _make_entry(
            transaction_type="DEBIT",
            date_posted=date(2026, 1, 15),
            amount=Decimal("-49.99"),
            description="Example Retailer Office Supplies",
            category="Software and IT",
            note="tax deductible",
        )
        out = tmp_path / "ledger.xlsx"
        export_ledger({"Jan2026_CAD": [entry]}, out, year=2026)

        wb = openpyxl.load_workbook(out)
        ws = wb["Jan2026_CAD"]
        # Row 2 is the first data row (row 1 = headers).
        row = [ws.cell(row=2, column=c).value for c in range(1, 9)]
        # Columns: AccountNumber, TxnType, DatePosted, Amount, Description,
        #          Category, Invoice/Cheque, Note
        assert row[1] == "DEBIT"
        assert row[3] == pytest.approx(-49.99)
        assert row[4] == "Example Retailer Office Supplies"
        assert row[5] == "Software and IT"
        assert row[7] == "tax deductible"
        wb.close()


# ---------------------------------------------------------------------------
# export_ledger — hyperlinks
# ---------------------------------------------------------------------------

class TestExportLedgerHyperlinks:
    def test_export_ledger_hyperlinks(self, tmp_path: Path):
        entry = _make_entry(
            document_link="202601 Jan/20260115-ExampleRetailerOfficeSupplies.pdf",
        )
        out = tmp_path / "ledger.xlsx"
        export_ledger({"Jan2026_CAD": [entry]}, out, year=2026)

        wb = openpyxl.load_workbook(out)
        ws = wb["Jan2026_CAD"]
        cell = ws.cell(row=2, column=7)  # Invoice/Cheque column (G)
        assert cell.hyperlink is not None
        wb.close()


# ---------------------------------------------------------------------------
# export_ledger — IndexSheet
# ---------------------------------------------------------------------------

class TestExportLedgerIndexSheet:
    def test_export_ledger_index_sheet(self, tmp_path: Path):
        out = tmp_path / "ledger.xlsx"
        entries = {
            "Jan2026_CAD": [_make_entry()],
            "Feb2026_CAD": [
                _make_entry(month="Feb2026", date_posted=date(2026, 2, 1)),
            ],
        }
        export_ledger(entries, out, year=2026)

        wb = openpyxl.load_workbook(out)
        assert "IndexSheet" in wb.sheetnames
        ws = wb["IndexSheet"]
        # IndexSheet should reference the monthly sheet names somewhere in its
        # data cells.
        cell_values = []
        for row in ws.iter_rows(values_only=True):
            cell_values.extend(v for v in row if v is not None)
        joined = " ".join(str(v) for v in cell_values)
        assert "Jan2026_CAD" in joined
        assert "Feb2026_CAD" in joined
        wb.close()


# ---------------------------------------------------------------------------
# export_ledger — NET INCOME sheet
# ---------------------------------------------------------------------------

class TestExportLedgerNetIncome:
    def test_export_ledger_net_income_sheet(self, tmp_path: Path):
        income_entry = _make_entry(
            transaction_type="CREDIT",
            amount=Decimal("5000.00"),
            category="Income",
            description="Client payment",
        )
        expense_entry = _make_entry(
            transaction_type="DEBIT",
            amount=Decimal("-200.00"),
            category="Software and IT",
            description="Keyboard",
        )
        out = tmp_path / "ledger.xlsx"
        export_ledger({"Jan2026_CAD": [income_entry, expense_entry]}, out, year=2026)

        wb = openpyxl.load_workbook(out)
        assert "NET INCOME" in wb.sheetnames
        wb.close()


# ---------------------------------------------------------------------------
# export_ledger — empty entries
# ---------------------------------------------------------------------------

class TestExportLedgerEmpty:
    def test_export_ledger_empty_entries(self, tmp_path: Path):
        out = tmp_path / "ledger.xlsx"
        export_ledger({}, out, year=2026)

        assert out.exists()
        wb = openpyxl.load_workbook(out)
        assert "IndexSheet" in wb.sheetnames
        assert "NET INCOME" in wb.sheetnames
        wb.close()


# ---------------------------------------------------------------------------
# export_ledger — multiple account types
# ---------------------------------------------------------------------------

class TestExportLedgerMultipleAccounts:
    def test_export_ledger_multiple_accounts(self, tmp_path: Path):
        cad_entry = _make_entry(account=AccountType.CAD, month="Jan2026")
        usd_entry = _make_entry(
            account=AccountType.USD,
            month="Jan2026",
            currency="USD",
            description="US vendor payment",
        )
        out = tmp_path / "ledger.xlsx"
        export_ledger(
            {
                "Jan2026_CAD": [cad_entry],
                "Jan2026_USD": [usd_entry],
            },
            out,
            year=2026,
        )
        wb = openpyxl.load_workbook(out)
        assert "Jan2026_CAD" in wb.sheetnames
        assert "Jan2026_USD" in wb.sheetnames
        wb.close()


# ---------------------------------------------------------------------------
# write_ledger_entries — append behaviour
# ---------------------------------------------------------------------------

class TestWriteLedgerEntriesAppends:
    def test_write_ledger_entries_appends(self, tmp_path: Path):
        out = tmp_path / "ledger.xlsx"
        wb = openpyxl.Workbook()

        entry_a = _make_entry(description="First purchase")
        entry_b = _make_entry(description="Second purchase")

        write_ledger_entries(wb, [entry_a], AccountType.CAD, 2026, 1)
        write_ledger_entries(wb, [entry_b], AccountType.CAD, 2026, 1)

        ws = wb["Jan2026_CAD"]
        # Row 1 = headers, rows 2 & 3 = data
        assert ws.cell(row=2, column=5).value == "First purchase"
        assert ws.cell(row=3, column=5).value == "Second purchase"
        wb.close()


# ---------------------------------------------------------------------------
# export_ledger_csv
# ---------------------------------------------------------------------------

class TestExportLedgerCSV:
    def test_csv_creates_file(self, tmp_path: Path):
        out = tmp_path / "ledger.csv"
        entries = {"Jan2026_CAD": [_make_entry()]}
        export_ledger_csv(entries, out)
        assert out.exists()

    def test_csv_has_header_row(self, tmp_path: Path):
        out = tmp_path / "ledger.csv"
        export_ledger_csv({"Jan2026_CAD": [_make_entry()]}, out)
        lines = out.read_text().strip().split("\n")
        header = lines[0]
        assert "Account" in header
        assert "Transaction Type" in header
        assert "Transaction Amount" in header
        assert "Description" in header
        assert "Category" in header

    def test_csv_contains_entry_data(self, tmp_path: Path):
        entry = _make_entry(
            amount=Decimal("-49.99"),
            description="Example Retailer Office Supplies",
            category="Software and IT",
        )
        out = tmp_path / "ledger.csv"
        export_ledger_csv({"Jan2026_CAD": [entry]}, out)
        content = out.read_text()
        assert "Example Retailer Office Supplies" in content
        assert "-49.99" in content
        assert "Software and IT" in content
        assert "CAD" in content

    def test_csv_multiple_accounts(self, tmp_path: Path):
        cad_entry = _make_entry(description="CAD purchase")
        usd_entry = _make_entry(
            account=AccountType.USD, description="USD purchase", currency="USD",
        )
        out = tmp_path / "ledger.csv"
        export_ledger_csv({
            "Jan2026_CAD": [cad_entry],
            "Jan2026_USD": [usd_entry],
        }, out)
        content = out.read_text()
        assert "CAD purchase" in content
        assert "USD purchase" in content

    def test_csv_date_format(self, tmp_path: Path):
        entry = _make_entry(date_posted=date(2026, 1, 15))
        out = tmp_path / "ledger.csv"
        export_ledger_csv({"Jan2026_CAD": [entry]}, out)
        content = out.read_text()
        assert "2026-01-15" in content

    def test_transactions_csv_date_format_yyyymmdd(self, tmp_path: Path):
        """Dates in the transactions CSV are YYYYMMDD integers, not ISO text."""
        txn = _make_transaction(date_posted=date(2026, 2, 4))
        out = tmp_path / "txns.csv"
        export_transactions_csv([txn], out)
        content = out.read_text()
        assert "20260204" in content

    def test_csv_empty_entries(self, tmp_path: Path):
        out = tmp_path / "ledger.csv"
        export_ledger_csv({}, out)
        assert out.exists()
        lines = out.read_text().strip().split("\n")
        assert len(lines) == 1  # header only


# ---------------------------------------------------------------------------
# export_transactions_csv — exports directly from transactions table
# ---------------------------------------------------------------------------

def _make_transaction(
    account: AccountType = AccountType.CAD,
    transaction_type: str = "DEBIT",
    date_posted: date = date(2026, 2, 4),
    amount: Decimal = Decimal("-1234.56"),
    currency: str = "CAD",
    description: str = "US$Transfer",
    txn_id: int = 1,
) -> Transaction:
    return Transaction(
        id=txn_id,
        account=account,
        transaction_type=transaction_type,
        date_posted=date_posted,
        amount=amount,
        currency=currency,
        description=description,
        source_file="test.csv",
        source_row=1,
        status=TransactionStatus.UNMATCHED,
    )


class TestExportTransactionsCSV:
    def test_creates_file(self, tmp_path: Path):
        txns = [_make_transaction()]
        out = tmp_path / "txns.csv"
        export_transactions_csv(txns, out)
        assert out.exists()

    def test_header_row(self, tmp_path: Path):
        out = tmp_path / "txns.csv"
        export_transactions_csv([_make_transaction()], out)
        header = out.read_text().split("\n")[0]
        assert "Account Number" in header
        assert "Transaction Amount" in header
        assert "Description" in header
        assert "Invoice/Cheque" in header
        assert "Category" in header
        assert "Note" in header

    def test_transaction_data_in_output(self, tmp_path: Path):
        txn = _make_transaction(
            description="IncomingWirePayment,INCOMINGWIRE",
            amount=Decimal("5000.00"),
            transaction_type="CREDIT",
        )
        out = tmp_path / "txns.csv"
        export_transactions_csv([txn], out)
        content = out.read_text()
        assert "IncomingWirePayment" in content
        assert "5000.00" in content
        assert "CREDIT" in content

    def test_unmatched_has_empty_invoice(self, tmp_path: Path):
        txn = _make_transaction(txn_id=99)
        out = tmp_path / "txns.csv"
        export_transactions_csv([txn], out, matched_docs={})
        lines = out.read_text().strip().split("\n")
        # Invoice/Cheque and Note should be empty (last two fields)
        data_line = lines[1]
        assert data_line.endswith(",,")  # empty Invoice/Cheque and Note

    def test_matched_has_document_path(self, tmp_path: Path):
        txn = _make_transaction(txn_id=42)
        out = tmp_path / "txns.csv"
        export_transactions_csv(
            [txn], out,
            matched_docs={42: "data/2026/202602 Feb/20260204-USTransfer.pdf"},
        )
        content = out.read_text()
        assert "20260204-USTransfer.pdf" in content

    def test_sorted_by_date(self, tmp_path: Path):
        txn_late = _make_transaction(date_posted=date(2026, 2, 20), txn_id=1)
        txn_early = _make_transaction(date_posted=date(2026, 2, 2), txn_id=2)
        out = tmp_path / "txns.csv"
        export_transactions_csv([txn_late, txn_early], out)
        lines = out.read_text().strip().split("\n")
        assert "20260202" in lines[1]  # early date first
        assert "20260220" in lines[2]  # late date second


# ---------------------------------------------------------------------------
# Database: get_latest_transaction_month
# ---------------------------------------------------------------------------

class TestGetLatestTransactionMonth:
    def test_returns_none_when_empty(self, tmp_path: Path):
        from db.database import Database
        db = Database(str(tmp_path / "test.db"))
        assert db.get_latest_transaction_month() is None

    def test_returns_latest_month(self, tmp_path: Path):
        from db.database import Database
        db = Database(str(tmp_path / "test.db"))
        txn = _make_transaction(date_posted=date(2026, 2, 15))
        db.insert_transaction(txn)
        result = db.get_latest_transaction_month()
        assert result == (2026, 2)
