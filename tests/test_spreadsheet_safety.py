"""core/spreadsheet_safety.py — document text must not become a live formula.

A vendor name or bank descriptor is written by whoever issued the document. If
it starts with ``=`` and reaches an export verbatim, the accountant's
spreadsheet executes it. These tests pin both halves of the guard, and the
property that matters most in a bookkeeping export: an amount is never touched.
"""

from __future__ import annotations

import csv
import io

import openpyxl
import pytest

from core.spreadsheet_safety import SafeCsvWriter, csv_safe, neutralize_workbook_formulas

HOSTILE = [
    '=HYPERLINK("http://evil.example.com","Invoice")',
    "=1+1",
    "+1+1",
    "-1+1",
    "@SUM(A1:A9)",
    "\t=1+1",
    "\r=1+1",
]


@pytest.mark.parametrize("text", HOSTILE)
def test_csv_safe_makes_a_formula_leader_inert(text: str) -> None:
    assert csv_safe(text) == "'" + text


@pytest.mark.parametrize("amount", ["-67.20", "+9600.00", "-1,234.56", "-4", "0.00", "-0.5"])
def test_csv_safe_never_alters_an_amount(amount: str) -> None:
    assert csv_safe(amount) == amount


@pytest.mark.parametrize("value", ["ORCHID TELECOM PREAUTH PMT", "", "Uncategorized", "a=b", None, 12, -3.5])
def test_csv_safe_leaves_ordinary_values_alone(value) -> None:
    assert csv_safe(value) == value


def test_safe_csv_writer_guards_text_and_keeps_numbers() -> None:
    out = io.StringIO()
    writer = SafeCsvWriter(csv.writer(out))
    writer.writerow(["2026-01-15", "=cmd|' /C calc'!A0", "-67.20", "CAD"])
    writer.writerows([["2026-01-16", "Bluepeak Software Inc", "-300.00", "CAD"]])
    rows = list(csv.reader(io.StringIO(out.getvalue())))
    assert rows[0] == ["2026-01-15", "'=cmd|' /C calc'!A0", "-67.20", "CAD"]
    assert rows[1] == ["2026-01-16", "Bluepeak Software Inc", "-300.00", "CAD"]


def test_workbook_formula_cells_are_saved_as_text_with_the_text_unchanged(tmp_path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = HOSTILE[0]
    ws["A2"] = "Cedarview Supplies"
    ws["A3"] = -67.2
    assert ws["A1"].data_type == "f"  # what openpyxl infers without the guard

    assert neutralize_workbook_formulas(wb) == 1

    path = tmp_path / "out.xlsx"
    wb.save(path)
    back = openpyxl.load_workbook(path)
    sheet = back.active
    assert sheet["A1"].data_type == "s"
    assert sheet["A1"].value == HOSTILE[0]
    assert sheet["A2"].value == "Cedarview Supplies"
    assert sheet["A3"].value == -67.2


def test_the_ledger_workbook_export_neutralises_a_hostile_description(tmp_path) -> None:
    """End to end through the real writer: no formula-typed cell survives."""
    from datetime import date
    from decimal import Decimal

    from core.ledger import export_ledger
    from models.ledger_entry import LedgerEntry
    from models.transaction import AccountType

    entry = LedgerEntry(
        account=AccountType.CAD,
        month="Jan2026",
        transaction_type="DEBIT",
        date_posted=date(2026, 1, 15),
        amount=Decimal("-67.20"),
        currency="CAD",
        description=HOSTILE[0],
        category="Utilities and telephone",
    )
    out = tmp_path / "ledger.xlsx"
    export_ledger({"Jan2026_CAD": [entry]}, out, year=2026)

    book = openpyxl.load_workbook(out)
    cells = [c for ws in book.worksheets for row in ws.iter_rows() for c in row]
    assert [c.coordinate for c in cells if c.data_type == "f"] == []
    assert HOSTILE[0] in {c.value for c in cells}, "the description must survive verbatim"


def test_the_qbo_csv_export_guards_a_hostile_description(tmp_path) -> None:
    from datetime import date
    from decimal import Decimal

    from core.ledger import export_qbo_csv
    from models.transaction import AccountType, Transaction

    txn = Transaction(
        id=1,
        account=AccountType.CAD,
        transaction_type="DEBIT",
        date_posted=date(2026, 1, 15),
        amount=Decimal("-67.20"),
        currency="CAD",
        description=HOSTILE[0],
        source_file="demo_bank_cad.csv",
        source_row=1,
    )
    out = tmp_path / "qbo.csv"
    export_qbo_csv([txn], out)
    rows = list(csv.reader(out.open(encoding="utf-8")))
    flat = [cell for row in rows[1:] for cell in row]
    assert "'" + HOSTILE[0] in flat
    assert not any(cell.startswith("=") for cell in flat)
    assert "-67.20" in flat, "the amount must leave untouched"
