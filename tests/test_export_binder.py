"""ZIP audit binder — every proof link resolves to a file inside the archive.

The binder stores each proof under ``proof_of_transaction/`` with a decorated
``{YYYYMMDD}-{Vendor}-{Amount}-{Currency}.pdf`` name, while the source document
is held under its storage key. Build the HTML link or the XLSX hyperlink from
the storage-key basename instead of the stored name and every link in the
binder points at nothing — an auditor opening it finds no proofs at all.

These tests pin link and entry to the same name: in the HTML index, in the
workbook's Proof of Transaction column, for a document shared by two
transactions, and for a fetch that failed (which must produce no link rather
than a broken one).
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date
from decimal import Decimal
from unittest.mock import patch

import pytest

from core.export_binder import _FILE_GONE
from models.transaction import AccountType, Transaction, TransactionStatus

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def _make_txn(
    txn_id: int,
    desc: str,
    amount: Decimal,
    currency: str = "CAD",
    txn_type: str = "DEBIT",
    status: TransactionStatus = TransactionStatus.MATCHED,
    acct: AccountType = AccountType.CAD,
    dt: date = date(2026, 1, 2),
) -> Transaction:
    txn = Transaction(
        account=acct,
        transaction_type=txn_type,
        date_posted=dt,
        amount=amount,
        currency=currency,
        description=desc,
        source_file="bank.csv",
        source_row=1,
        status=status,
    )
    txn.id = txn_id
    return txn


SAMPLE_PDF_BYTES = b"%PDF-1.4 dummy content for unit test"


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestAuditBinderReceiptLinks:
    """Verify that HTML links in the binder match actual ZIP file entries."""

    @pytest.mark.asyncio
    async def test_html_links_match_zip_entries(self):
        """Every proof-of-transaction link in the HTML must correspond to an
        actual file inside proof_of_transaction/ in the ZIP."""
        from core.export_binder import generate_audit_binder

        txns = [
            _make_txn(1, "PRE-AUTHORIZED PAYMENT", Decimal("24.00"), dt=date(2026, 1, 2)),
            _make_txn(2, "ORCHID TELECOM", Decimal("88.00"), dt=date(2026, 1, 5)),
            _make_txn(3, "BLUEPEAK SOFTWARE INC", Decimal("300.00"), dt=date(2026, 1, 6)),
            _make_txn(4, "E-TRANSFER", Decimal("500.00"), status=TransactionStatus.IGNORED, dt=date(2026, 1, 7)),
            _make_txn(5, "EXAMPLE RETAILER", Decimal("45.00"), status=TransactionStatus.UNMATCHED, dt=date(2026, 1, 8)),
        ]

        doc_paths = {
            1: "receipts/0123456789ab_statement-export.pdf",
            2: "receipts/ab12cd34ef56_telecom-bill.pdf",
            3: "receipts/ff99ee88dd77_software-invoice.pdf",
            # 4 is IGNORED (no doc), 5 is UNMATCHED (no doc)
        }

        vendors = {
            1: "Example Retailer",
            2: "Orchid Telecom",
            3: "Bluepeak Software Inc",
        }

        # Mock the file store — every fetch succeeds, returning dummy PDF bytes
        async def mock_fetch(path: str, user_id: str) -> tuple[bytes | None, str | None]:
            return SAMPLE_PDF_BYTES, None

        with patch("core.export_binder._fetch_receipt_bytes", side_effect=mock_fetch):
            zip_bytes = await generate_audit_binder(
                company_name="Test Corp",
                year=2026,
                month=1,
                label="Jan2026",
                transactions=txns,
                doc_paths=doc_paths,
                vendors=vendors,
                user_id="test-user",
            )

        # Parse the ZIP
        zf = zipfile.ZipFile(io.BytesIO(zip_bytes), "r")
        zip_names = set(zf.namelist())

        # Verify proof_of_transaction/ files exist in ZIP
        proof_files_in_zip = {
            name for name in zip_names if name.startswith("proof_of_transaction/")
        }
        assert len(proof_files_in_zip) == 3, (
            f"Expected 3 proof files in ZIP (for 3 matched txns), got {len(proof_files_in_zip)}: {proof_files_in_zip}"
        )

        # Verify filenames follow the standardized format
        for pf in proof_files_in_zip:
            basename = pf.replace("proof_of_transaction/", "")
            # Pattern: YYYYMMDD-Vendor-Amount-Currency.pdf
            assert re.match(r"^\d{8}-.+-.+\..+$", basename), (
                f"Proof file '{basename}' does not follow YYYYMMDD-Vendor-Amount-Currency.ext format"
            )

        # Verify index.html exists
        assert "index.html" in zip_names, "index.html must exist in ZIP"

        # Parse HTML and extract all proof-of-transaction links
        html_content = zf.read("index.html").decode("utf-8")
        href_pattern = re.compile(r'href="(proof_of_transaction/[^"]+)"')
        html_links = set(href_pattern.findall(html_content))

        assert len(html_links) > 0, "HTML should contain at least one proof-of-transaction link"

        # THE KEY CHECK: every link in the HTML must point to an actual file in the ZIP
        for link in html_links:
            # URL-decode percent-encoded characters for comparison
            from urllib.parse import unquote
            decoded_link = unquote(link)
            assert decoded_link in zip_names, (
                f"HTML link '{decoded_link}' does not match any ZIP entry.\n"
                f"ZIP entries: {sorted(proof_files_in_zip)}"
            )

        zf.close()

    @pytest.mark.asyncio
    async def test_failed_fetch_does_not_produce_broken_link(self):
        """If a file fetch fails, the HTML should NOT contain a link for it."""
        from core.export_binder import generate_audit_binder

        txns = [
            _make_txn(1, "EXAMPLE RETAILER PAYMENT", Decimal("26.50"), dt=date(2026, 1, 2)),
            _make_txn(2, "ORCHID TELECOM PAYMENT", Decimal("88.00"), dt=date(2026, 1, 5)),
        ]

        doc_paths = {
            1: "receipts/aaa_retailer.pdf",    # This will FAIL to fetch
            2: "receipts/bbb_telecom.pdf",     # This will succeed
        }

        vendors = {1: "Example Retailer", 2: "Orchid Telecom"}

        async def mock_fetch_partial(
            path: str, user_id: str
        ) -> tuple[bytes | None, str | None]:
            if "aaa_retailer" in path:
                return None, _FILE_GONE  # the file store could not produce it
            return SAMPLE_PDF_BYTES, None

        with patch("core.export_binder._fetch_receipt_bytes", side_effect=mock_fetch_partial):
            zip_bytes = await generate_audit_binder(
                company_name="Test Corp",
                year=2026,
                month=1,
                label="Jan2026",
                transactions=txns,
                doc_paths=doc_paths,
                vendors=vendors,
                user_id="test-user",
            )

        zf = zipfile.ZipFile(io.BytesIO(zip_bytes), "r")
        zip_names = set(zf.namelist())

        # Only 1 proof file should be in ZIP (the one that succeeded)
        proof_files = {n for n in zip_names if n.startswith("proof_of_transaction/")}
        assert len(proof_files) == 1, f"Expected 1 proof file, got {len(proof_files)}: {proof_files}"

        # HTML should not link to the failed file
        html_content = zf.read("index.html").decode("utf-8")
        href_pattern = re.compile(r'href="(proof_of_transaction/[^"]+)"')
        html_links = set(href_pattern.findall(html_content))

        # All HTML links must point to files in the ZIP
        from urllib.parse import unquote
        for link in html_links:
            decoded = unquote(link)
            assert decoded in zip_names, (
                f"HTML link '{decoded}' points to non-existent ZIP entry (broken link!)"
            )

        # The failed-fetch transaction should show mdash, not a clickable link
        assert "Retailer" not in " ".join(html_links), (
            "Example Retailer transaction link should not appear since its file fetch failed"
        )

        zf.close()

    @pytest.mark.asyncio
    async def test_xlsx_proof_column_matches_zip(self):
        """XLSX 'Proof of Transaction' column should reference files that exist in ZIP."""
        from core.export_binder import generate_audit_binder

        txns = [
            _make_txn(1, "ORCHID TELECOM PAYMENT", Decimal("88.00"), dt=date(2026, 1, 5)),
        ]
        doc_paths = {1: "receipts/xyz_telecom.pdf"}
        vendors = {1: "Orchid Telecom"}

        async def mock_fetch(path: str, user_id: str) -> tuple[bytes | None, str | None]:
            return SAMPLE_PDF_BYTES, None

        with patch("core.export_binder._fetch_receipt_bytes", side_effect=mock_fetch):
            zip_bytes = await generate_audit_binder(
                company_name="Test Corp",
                year=2026,
                month=1,
                label="Jan2026",
                transactions=txns,
                doc_paths=doc_paths,
                vendors=vendors,
                user_id="test-user",
            )

        zf = zipfile.ZipFile(io.BytesIO(zip_bytes), "r")
        zip_names = set(zf.namelist())
        proof_files = {n.replace("proof_of_transaction/", "") for n in zip_names if n.startswith("proof_of_transaction/")}

        # Read XLSX and check proof column
        import openpyxl
        xlsx_names = [n for n in zip_names if n.endswith(".xlsx")]
        assert len(xlsx_names) == 1, f"Expected 1 XLSX, got {xlsx_names}"

        xlsx_bytes = zf.read(xlsx_names[0])
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

        # Check the CAD sheet (title block rows 1-2, header row 3, data from row 4)
        for sheet_name in wb.sheetnames:
            if "CAD" in sheet_name:
                ws = wb[sheet_name]
                # Row 4 is first data row; column 10 is "Proof of Transaction"
                assert ws.cell(row=3, column=10).value == "Proof of Transaction"
                proof_cell = ws.cell(row=4, column=10)
                proof_value = proof_cell.value
                assert proof_value is not None, "Proof of Transaction column should have a value"
                # The cell must be a clickable RELATIVE hyperlink to the
                # archived file so users can open the proof directly from
                # the extracted ZIP folder.
                assert proof_cell.hyperlink is not None, "proof cell should be a hyperlink"
                assert proof_cell.hyperlink.target in zip_names, (
                    f"hyperlink target '{proof_cell.hyperlink.target}' is not a ZIP entry"
                )
                assert proof_value in proof_files, (
                    f"XLSX proof '{proof_value}' not found in ZIP proof_of_transaction/ folder.\n"
                    f"Available: {proof_files}"
                )
                break
        else:
            pytest.fail("No CAD sheet found in XLSX")

        wb.close()
        zf.close()

    @pytest.mark.asyncio
    async def test_multiple_txns_same_document(self):
        """Multiple transactions sharing one document should all link to same file."""
        from core.export_binder import generate_audit_binder

        txns = [
            _make_txn(1, "Payment Part 1", Decimal("50.00"), dt=date(2026, 1, 2)),
            _make_txn(2, "Payment Part 2", Decimal("50.00"), dt=date(2026, 1, 3)),
        ]

        # Both transactions point to the same receipt file
        shared_path = "receipts/shared_invoice.pdf"
        doc_paths = {1: shared_path, 2: shared_path}
        vendors = {1: "SharedVendor", 2: "SharedVendor"}

        async def mock_fetch(path: str, user_id: str) -> tuple[bytes | None, str | None]:
            return SAMPLE_PDF_BYTES, None

        with patch("core.export_binder._fetch_receipt_bytes", side_effect=mock_fetch):
            zip_bytes = await generate_audit_binder(
                company_name="Test Corp",
                year=2026,
                month=1,
                label="Jan2026",
                transactions=txns,
                doc_paths=doc_paths,
                vendors=vendors,
                user_id="test-user",
            )

        zf = zipfile.ZipFile(io.BytesIO(zip_bytes), "r")
        zip_names = set(zf.namelist())
        proof_files = [n for n in zip_names if n.startswith("proof_of_transaction/")]

        # Only one file should be written (shared), not two copies
        assert len(proof_files) == 1, (
            f"Expected 1 proof file (shared doc), got {len(proof_files)}: {proof_files}"
        )

        # HTML should have 2 links, both pointing to the same file
        html_content = zf.read("index.html").decode("utf-8")
        href_pattern = re.compile(r'href="(proof_of_transaction/[^"]+)"')
        html_links = href_pattern.findall(html_content)
        assert len(html_links) == 2, f"Expected 2 links (one per txn), got {len(html_links)}"

        from urllib.parse import unquote
        decoded_links = {unquote(link) for link in html_links}
        assert len(decoded_links) == 1, "Both links should point to the same file"
        assert decoded_links.pop() in zip_names, "Shared link must exist in ZIP"

        zf.close()

    @pytest.mark.asyncio
    async def test_zip_uses_proof_of_transaction_folder(self):
        """Archived proofs live under proof_of_transaction/, never receipts/.

        ``receipts/`` is the storage-key prefix of the source document. An
        archive that reused it would collide with the storage layout and make
        the binder's own links ambiguous.
        """
        from core.export_binder import generate_audit_binder

        txns = [_make_txn(1, "Test Payment", Decimal("10.00"))]
        doc_paths = {1: "receipts/test.pdf"}
        vendors = {1: "TestVendor"}

        async def mock_fetch(path: str, user_id: str) -> tuple[bytes | None, str | None]:
            return SAMPLE_PDF_BYTES, None

        with patch("core.export_binder._fetch_receipt_bytes", side_effect=mock_fetch):
            zip_bytes = await generate_audit_binder(
                company_name="Test Corp",
                year=2026,
                month=1,
                label="Jan2026",
                transactions=txns,
                doc_paths=doc_paths,
                vendors=vendors,
                user_id="test-user",
            )

        zf = zipfile.ZipFile(io.BytesIO(zip_bytes), "r")
        zip_names = zf.namelist()

        # No entries should use the storage-key "receipts/" prefix
        receipts_entries = [n for n in zip_names if n.startswith("receipts/")]
        assert len(receipts_entries) == 0, (
            f"ZIP should not contain a 'receipts/' folder. Found: {receipts_entries}"
        )

        # Must use proof_of_transaction/
        proof_entries = [n for n in zip_names if n.startswith("proof_of_transaction/")]
        assert len(proof_entries) > 0, "ZIP must contain proof_of_transaction/ folder"

        # HTML must also use proof_of_transaction/, not receipts/
        html_content = zf.read("index.html").decode("utf-8")
        assert 'href="receipts/' not in html_content, (
            "HTML links must use proof_of_transaction/, not receipts/"
        )
        assert 'href="proof_of_transaction/' in html_content, (
            "HTML links must use proof_of_transaction/ prefix"
        )

        zf.close()
