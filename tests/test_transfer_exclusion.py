"""Internal-transfer exclusion from income/expense aggregations.

Money moved between the company's own accounts is not income and not an
expense. Counted naively, a USD<->CAD internal transfer (or a credit-card
payment) books as income on the receiving account and as an expense on the
sending one, which double-counts it in the binder XLSX, the ledger NET INCOME
sheet and the PDF/HTML report totals.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import openpyxl

from core.transfer_exclusion import (
    TRANSFER_LINK_TYPES,
    is_transfer_category,
    normalize_category,
    transfer_exclusion_sql,
)
from models.transaction import AccountType, Transaction, TransactionStatus


def _make_txn(
    txn_id: int,
    desc: str,
    amount: Decimal,
    txn_type: str = "DEBIT",
    category: str | None = None,
    acct: AccountType = AccountType.CAD,
    dt: date = date(2026, 1, 2),
) -> Transaction:
    txn = Transaction(
        account=acct,
        transaction_type=txn_type,
        date_posted=dt,
        amount=amount,
        currency="CAD",
        description=desc,
        source_file="bank.csv",
        source_row=1,
        status=TransactionStatus.UNMATCHED,
        category=category,
    )
    txn.id = txn_id
    return txn


class TestIsTransferCategory:
    def test_known_transfer_names(self):
        for name in [
            "Credit Card Payment",
            "Internal Transfer",
            "Inter-Account Transfers",  # a plural/hyphenated variant a user may type
            "USD to CAD",
            "CAD to USD",
        ]:
            assert is_transfer_category(name), name

    def test_case_and_punctuation_insensitive(self):
        assert is_transfer_category("inter account transfer")
        assert is_transfer_category("USD-to-CAD")
        assert is_transfer_category("credit_card_payment")

    def test_real_pl_categories_not_excluded(self):
        for name in [
            "Bank Fee",       # real expense — must stay in totals
            "Interest",       # real income — must stay in totals
            "Contract Fee",
            "Income",
            "Refund/Credit",
            "Travel",
            # Wages/payroll are P&L activity (the rule pre-pass assigns the
            # category to every cheque), so excluding them would understate
            # expenses.
            "Wage Payment",
            "Payroll Deposit",
            # Owner's draws count as money out — a missing invoice is a receipt
            # problem, not grounds to drop the cashflow.
            "Owner's Draw & Shareholder Distributions",
            "Owner's Draw",
            None,
            "",
        ]:
            assert not is_transfer_category(name), name

    def test_normalize(self):
        assert normalize_category("Inter-Account Transfers") == "interaccounttransfers"
        assert normalize_category(None) == ""


class TestTransferExclusionSql:
    def test_predicate_shape(self):
        sql, params = transfer_exclusion_sql(alias="t")
        assert "t.category" in sql
        assert "transaction_links" in sql
        assert "user_categories" in sql
        assert sql.count("%s") == len(params) == 2
        assert params[1] == list(TRANSFER_LINK_TYPES)
        assert "REFUND" not in params[1]


class TestBinderXlsxTotals:
    def _summary_totals(self, xlsx_bytes: bytes) -> tuple[float, float, float]:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb["Summary"]
        # Find the grand-total row. Totals are per currency ("TOTAL (CAD)");
        # these fixtures are single-currency so exactly one row exists.
        # Layout: Account | Currency | Transactions | Income | Expenses | Net
        for row in ws.iter_rows(min_col=1, max_col=6):
            if isinstance(row[0].value, str) and row[0].value.startswith("TOTAL"):
                return (
                    float(row[3].value or 0),
                    float(row[4].value or 0),
                    float(row[5].value or 0),
                )
        raise AssertionError("TOTAL row not found in Summary sheet")

    def test_transfer_pair_excluded_from_totals(self):
        from core.export_binder import _build_xlsx

        txns = [
            # Real business activity
            _make_txn(1, "CLIENT WIRE", Decimal("5000.00"), txn_type="CREDIT",
                      category="Income"),
            _make_txn(2, "ORCHID TELECOM BILL", Decimal("-100.00"), txn_type="DEBIT",
                      category="Office IT"),
            # Internal transfer pair: USD -> CAD, marked by category on one
            # side and by transfer link (ids set) on the other.
            _make_txn(3, "USD TFR TO CAD", Decimal("1300.00"), txn_type="CREDIT",
                      category="USD to CAD"),
            _make_txn(4, "TRANSFER OUT", Decimal("-1000.00"), txn_type="DEBIT",
                      acct=AccountType.USD),
        ]

        xlsx = _build_xlsx(
            "TestCo", 2026, 1, "Jan2026", txns, {}, {},
            transfer_txn_ids={4},
        )
        income, expenses, net = self._summary_totals(xlsx)
        assert income == 5000.00
        assert expenses == 100.00
        assert net == 4900.00

    def test_without_exclusion_signals_totals_include_everything(self):
        from core.export_binder import _build_xlsx

        txns = [
            _make_txn(1, "CLIENT WIRE", Decimal("5000.00"), txn_type="CREDIT",
                      category="Income"),
            _make_txn(2, "ORCHID TELECOM BILL", Decimal("-100.00"), txn_type="DEBIT",
                      category="Office IT"),
        ]
        xlsx = _build_xlsx("TestCo", 2026, 1, "Jan2026", txns, {}, {})
        income, expenses, net = self._summary_totals(xlsx)
        assert income == 5000.00
        assert expenses == 100.00
        assert net == 4900.00


class TestLedgerNetIncomeTotals:
    def test_monthly_totals_exclude_transfers(self, tmp_path):
        from core.ledger import export_ledger
        from models.ledger_entry import LedgerEntry

        def entry(amount: str, txn_type: str, category: str | None) -> LedgerEntry:
            return LedgerEntry(
                account=AccountType.CAD,
                month="Jan2026",
                transaction_type=txn_type,
                date_posted=date(2026, 1, 5),
                amount=Decimal(amount),
                currency="CAD",
                description="x",
                category=category,
            )

        entries = {
            "Jan2026_CAD": [
                entry("5000.00", "CREDIT", "Income"),
                entry("-100.00", "DEBIT", "Office IT"),
                entry("1300.00", "CREDIT", "USD to CAD"),
                entry("-1300.00", "DEBIT", "Credit Card Payment"),
            ]
        }
        out = tmp_path / "ledger.xlsx"
        export_ledger(entries, out, year=2026)

        wb = openpyxl.load_workbook(out)
        ws = wb["NET INCOME"]
        # Row 2: Jan2026_CAD — income excl. transfer, expenses excl. transfer.
        # Expenses are positive magnitudes and Net = Income − Expenses, so the
        # result is correct regardless of the stored sign convention (legacy
        # CSV parsers store DEBITs negative; the LLM parser stores positives).
        # Totals are reported per currency (column 2), so a Wise tab holding
        # CAD and USD rows never sums the two into a number that is neither.
        assert ws.cell(row=1, column=2).value == "Currency"
        assert ws.cell(row=2, column=2).value == "CAD"
        assert float(ws.cell(row=2, column=3).value) == 5000.00
        assert float(ws.cell(row=2, column=4).value) == 100.00
        assert float(ws.cell(row=2, column=5).value) == 4900.00
